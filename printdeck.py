"""Print Deck: a real, editable Word document (.docx) with every card of a deck at the exact
physical Magic card size (63 x 88 mm), nine to an A4 page, for printing at 100% and cutting
out to sleeve - for personal use (playtesting your own decks), never for selling proxies.

* The exact printing of each card (the deck's printing / the build's chosen printing).
* Scryfall's 'large' image (672 x 936 px, about 270 dpi at card size) from the card viewer's
  cache (data/img_xl), otherwise downloaded once - politely, in the background, with progress.
* Fixed-layout Word tables: 3 x 3 cells of exactly 63 x 88 mm, no cell padding, rows that can't
  split or grow, so Word never rescales or moves a card. One table per page.
* Optional light-grey cutting guides on the cell edges (the card border, never the artwork).
* Double-faced cards: the front by default; optionally the backs on their own pages at the end.
* A card whose image can't be found is NOT replaced by another printing: its slot says so and
  it is listed in the result.
Nothing in the deck or collection is changed."""
import os
import time
from dataclasses import dataclass, field

import requests

import card_viewer
import deckrules

CARD_W_MM, CARD_H_MM = 63.0, 88.0
PAPER = {"A4": (210.0, 297.0)}
COLS, ROWS = 3, 3
PER_PAGE = COLS * ROWS
MIN_PRINTER_MARGIN_MM = 4.0   # most home printers can't print closer to the edge than this
GUIDE_COLOUR = "BFBFBF"        # light grey
GUIDE_SIZE = 2                 # eighths of a point (0.25 pt)
HEADERS = {"User-Agent": "HomeCardScanner/0.8 (personal deck printing)", "Accept": "image/*"}
NOTICE = ("Your printable deck is ready. Open the Word document and print on A4 paper at 100% / Actual Size. "
          "Do not use Fit to Page.")


@dataclass
class Options:
    include_commander: bool = True
    include_basics: bool = True
    include_backs: bool = False
    guides: bool = True
    paper: str = "A4"


@dataclass
class Item:
    name: str
    scryfall_id: str
    image_normal: str = ""
    image_small: str = ""
    role: str = "main"
    qty: int = 1
    double_faced: bool = False


@dataclass
class Plan:
    fronts: list = field(default_factory=list)   # [(Item, "front")] - one entry per physical card
    backs: list = field(default_factory=list)    # [(Item, "back")]

    @property
    def cards(self):
        return len(self.fronts)

    @property
    def pages(self):
        return pages_for(len(self.fronts)) + pages_for(len(self.backs))


def pages_for(n):
    return (n + PER_PAGE - 1) // PER_PAGE


def layout_check(paper="A4"):
    """(fits, margins (left/right, top/bottom) in mm, message). The cards are never scaled."""
    pw, ph = PAPER[paper]
    side = (pw - COLS * CARD_W_MM) / 2
    top = (ph - ROWS * CARD_H_MM) / 2
    if min(side, top) < MIN_PRINTER_MARGIN_MM:
        return False, (side, top), (f"3 x 3 cards need {COLS * CARD_W_MM:.0f} x {ROWS * CARD_H_MM:.0f} mm; on {paper} "
                                    f"that leaves only {min(side, top):.1f} mm of margin, less than most printers "
                                    "can print. Cards would not be scaled - choose another paper size.")
    return True, (side, top), ""


# ---- what to print ---------------------------------------------------------------------------

def is_double_faced(card):
    return " // " in (card.get("name") or "") and (card.get("layout") or "") not in ("split", "adventure",
                                                                                     "flip", "aftermath")


def items_from_deck(store, deck_id):
    """The deck's own rows - its exact printings and quantities (commander, partner and the 99)."""
    out = []
    for r in store.deck_cards(deck_id):
        if r["role"] not in ("commander", "partner", "main") or not r.get("scryfall_id"):
            continue
        card = store.card(r["scryfall_id"]) or {}
        out.append(Item(r["card_name"], r["scryfall_id"], r.get("image_normal") or card.get("image_normal") or "",
                        r.get("image_small") or card.get("image_small") or "", r["role"], int(r["quantity"] or 1),
                        is_double_faced(dict(card, name=r["card_name"]))))
    return out


def items_from_build(store, result):
    """A Build Deck result - the printings the builder chose (what Save Deck would use)."""
    out = []
    cmd = store.default_printing(result["commander"]) or {}
    if cmd:
        out.append(Item(cmd["name"], cmd["id"], cmd.get("image_normal") or "", cmd.get("image_small") or "",
                        "commander", 1, is_double_faced(cmd)))
    for c in result["cards"]:
        card = (store.card(c["scryfall_id"]) if c.get("scryfall_id") else None) or store.default_printing(c["name"])
        if not card:
            continue
        out.append(Item(c["name"], card["id"], card.get("image_normal") or "", card.get("image_small") or "", "main",
                        int(c.get("qty", 1)), is_double_faced(card)))
    return out


def plan(items, options=None):
    """Expand quantities into physical cards (a quantity of 4 prints four copies)."""
    o = options or Options()
    p = Plan()
    for it in items:
        if it.role in ("commander", "partner") and not o.include_commander:
            continue
        if deckrules.is_basic(it.name) and not o.include_basics:
            continue
        for _ in range(max(0, it.qty)):
            p.fronts.append((it, "front"))
            if o.include_backs and it.double_faced:
                p.backs.append((it, "back"))
    return p


# ---- images ----------------------------------------------------------------------------------

def image_path(it, face, session=None, timeout=25):
    """The cached large image of this exact printing (and face), downloading it if needed.
    Returns the path, or None when it isn't available - never another printing's picture."""
    path = card_viewer.xl_path(it.scryfall_id, face)
    if os.path.exists(path) and os.path.getsize(path) > 1000:
        return path
    url = card_viewer.large_url(dict(image_normal=it.image_normal, image_small=it.image_small), face)
    if not url:
        return None
    try:
        r = (session or requests).get(url, headers=HEADERS, timeout=timeout)
    except requests.RequestException:
        return None
    if r.status_code != 200 or r.content[:2] != b"\xff\xd8":
        return None
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path + ".part", "wb") as f:
        f.write(r.content)
    os.replace(path + ".part", path)
    time.sleep(0.05)  # polite to Scryfall's image servers
    return path


def has_back(it, session=None, timeout=15):
    """Is this printing really two-faced (a transform / modal double-faced card)? Split, Room,
    adventure and flip cards also have ' // ' in the name but only ONE printed face - Scryfall
    only has a back image for the real double-faced ones."""
    if os.path.exists(card_viewer.xl_path(it.scryfall_id, "back")):
        return True
    url = card_viewer.large_url(dict(image_normal=it.image_normal, image_small=it.image_small), "back")
    if not url:
        return False
    try:
        r = (session or requests).head(url, headers=HEADERS, timeout=timeout, allow_redirects=True)
        return r.status_code == 200
    except requests.RequestException:
        return True  # can't tell offline: try, and report it if the image is missing


# ---- the Word document -----------------------------------------------------------------------

def export(items, out_path, options=None, progress=None, fetch=image_path):
    """Write the .docx. Returns dict(path, cards, back_faces, pages, missing=[names], warnings)."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_BREAK, WD_LINE_SPACING
    from docx.shared import Mm, Pt

    o = options or Options()
    say = progress or (lambda f, t: None)
    fits, (side, top), msg = layout_check(o.paper)
    if not fits:
        raise ValueError(msg)
    p = plan(items, o)
    if not p.fronts:
        raise ValueError("Nothing to print with these options")
    session = requests.Session()
    if p.backs:
        say(0.02, "Checking which cards have a printed back face...")
        real = {it.scryfall_id: (has_back(it, session) if fetch is image_path else True) for it, _ in p.backs}
        p.backs = [(it, f) for it, f in p.backs if real[it.scryfall_id]]
    entries = p.fronts + p.backs
    # 1. images (once per printing + face, however many copies)
    paths, missing = {}, []
    unique = list(dict.fromkeys((it.scryfall_id, face) for it, face in entries))
    by_key = {(it.scryfall_id, face): it for it, face in entries}
    for i, key in enumerate(unique):
        it = by_key[key]
        say(0.05 + 0.75 * i / max(1, len(unique)), f"Getting card images {i + 1}/{len(unique)}: {it.name}")
        path = fetch(it, key[1], session) if fetch is image_path else fetch(it, key[1])
        paths[key] = path
        if path is None:
            missing.append(it.name + (" (back face)" if key[1] == "back" else ""))
    # 2. document
    say(0.85, "Laying out the pages...")
    doc = Document()
    sec = doc.sections[0]
    pw, ph = PAPER[o.paper]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Mm(pw), Mm(ph)
    # a little spare height for the one tiny paragraph Word needs after each table, so it never
    # spills onto an extra blank page
    sec.left_margin = sec.right_margin = Mm(side)
    sec.top_margin = sec.bottom_margin = Mm(top - 2)
    sec.header_distance = sec.footer_distance = Mm(0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(1)
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    doc.core_properties.title = "Printable deck (personal use)"
    doc.core_properties.comments = ("Card images from Scryfall; Magic: The Gathering card images are (c) Wizards of "
                                    "the Coast. For personal use only - not for sale.")
    body = doc.element.body
    for el in list(body):  # remove the default empty paragraph (keep the section properties)
        if el.tag.endswith("}p"):
            body.remove(el)

    def page_tables(seq):
        return [seq[i:i + PER_PAGE] for i in range(0, len(seq), PER_PAGE)]

    pages = page_tables(p.fronts) + page_tables(p.backs)
    for n, chunk in enumerate(pages):
        say(0.85 + 0.13 * n / len(pages), f"Page {n + 1} of {len(pages)}")
        table = doc.add_table(rows=ROWS, cols=COLS)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _fixed_layout(table, o.guides)
        for r, row in enumerate(table.rows):
            row.height = Mm(CARD_H_MM)
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            _cant_split(row)
            for c, cell in enumerate(row.cells):
                cell.width = Mm(CARD_W_MM)
                para = cell.paragraphs[0]
                para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(0)
                k = r * COLS + c
                if k >= len(chunk):
                    _no_borders(cell)  # an empty slot on the last page: no cutting lines
                    continue
                it, face = chunk[k]
                path = paths.get((it.scryfall_id, face))
                if path:
                    para.add_run().add_picture(path, width=Mm(CARD_W_MM), height=Mm(CARD_H_MM))
                else:
                    para.paragraph_format.space_before = Pt(80)
                    run = para.add_run(f"IMAGE MISSING\n{it.name}\n({face})")
                    run.font.size = Pt(10)
        last = n == len(pages) - 1
        tail = doc.add_paragraph()
        tail.paragraph_format.space_before = tail.paragraph_format.space_after = Pt(0)
        if not last:
            tail.add_run().add_break(WD_BREAK.PAGE)
    say(0.99, "Saving...")
    doc.save(out_path)
    warnings = []
    if missing:
        warnings.append(f"{len(missing)} card image(s) could not be found - their slots say IMAGE MISSING: "
                        + ", ".join(sorted(set(missing))[:12]))
    return dict(path=out_path, cards=p.cards, back_faces=len(p.backs), pages=len(pages), missing=missing,
                warnings=warnings)


def _set(el, tag, **attrs):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    node = el.find(qn(tag))
    if node is None:
        node = OxmlElement(tag)
        el.append(node)
    for k, v in attrs.items():
        node.set(qn(k), str(v))
    return node


def _fixed_layout(table, guides):
    """Fixed column widths, no cell padding, no spacing - and light-grey guides on the edges."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm
    tbl = table._tbl
    pr = tbl.tblPr
    _set(pr, "w:tblLayout", **{"w:type": "fixed"})
    _set(pr, "w:tblW", **{"w:w": str(int(Mm(CARD_W_MM * COLS).twips)), "w:type": "dxa"})
    mar = _set(pr, "w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        _set(mar, f"w:{side}", **{"w:w": "0", "w:type": "dxa"})
    _set(pr, "w:tblCellSpacing", **{"w:w": "0", "w:type": "dxa"})
    borders = _set(pr, "w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if guides:
            _set(borders, f"w:{side}", **{"w:val": "single", "w:sz": str(GUIDE_SIZE), "w:space": "0",
                                          "w:color": GUIDE_COLOUR})
        else:
            _set(borders, f"w:{side}", **{"w:val": "nil"})
    grid = tbl.tblGrid
    for col in grid.findall(qn("w:gridCol")):
        col.set(qn("w:w"), str(int(Mm(CARD_W_MM).twips)))
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(1, grid)


def _cant_split(row):
    trpr = row._tr.get_or_add_trPr()
    _set(trpr, "w:cantSplit")


def _no_borders(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    b = _set(tcpr, "w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        _set(b, f"w:{side}", **{"w:val": "nil"})


def default_filename(deck_name):
    safe = "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in (deck_name or "Deck")).strip() or "Deck"
    return f"{safe.replace(' ', '_')}_Printable.docx"
